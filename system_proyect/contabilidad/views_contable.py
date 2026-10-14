# <--- hecho por claude code (23-sep-2026): vistas del núcleo contable (catálogo, partidas, mayor, reportes,
# períodos y centros de costo). Las vistas NO calculan: llaman a contable.py. Autorización por permiso
# (ver_cuentas / administrar_cuentas / ver_partidas / crear_partidas / modificar_partidas / ver_reportes /
# administrar_configuracion); el superusuario siempre pasa.
import datetime
import io

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.db.models import Q, Count, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.views.decorators.http import require_POST

from .utils import puede, es_admin_contabilidad
from .models import (CuentaContable, CentroCosto, PeriodoContable, Asiento, DetalleAsiento,
                     TipoCuenta, TipoAsiento, EstadoAsiento)
from .forms import CuentaContableForm, CentroCostoForm, AsientoHeaderForm
from . import contable


def _gate(user, codename):
    if not puede(user, codename):
        raise PermissionDenied('No tienes permiso para esta acción contable.')


def _fecha(txt, default=None):
    try:
        return datetime.date.fromisoformat(txt) if txt else default
    except ValueError:
        return default


def _rango(request):
    """(desde, hasta) del GET; por defecto el mes actual."""
    hoy = timezone.now().date()
    desde = _fecha(request.GET.get('desde'), hoy.replace(day=1))
    hasta = _fecha(request.GET.get('hasta'), hoy)
    return desde, hasta


def _centro(request):
    cid = request.GET.get('centro')
    return CentroCosto.objects.filter(pk=cid).first() if cid else None


def _ctx(request, **extra):
    u = request.user
    base = {'nav_home_url': '/', 'es_admin': es_admin_contabilidad(u),
            'p_admin_cuentas': puede(u, 'administrar_cuentas'), 'p_crear': puede(u, 'crear_partidas'),
            'p_modificar': puede(u, 'modificar_partidas'), 'p_reportes': puede(u, 'ver_reportes'),
            'centros': CentroCosto.objects.filter(activo=True)}
    base.update(extra)
    return base


# ─────────────────────────── Catálogo de cuentas ───────────────────────────
@login_required
def cuentas_lista(request):
    _gate(request.user, 'ver_cuentas')
    q = (request.GET.get('q') or '').strip()
    tipo = request.GET.get('tipo') or ''
    cuentas = CuentaContable.objects.all()
    if q:
        cuentas = cuentas.filter(Q(codigo__icontains=q) | Q(nombre__icontains=q))
    if tipo:
        cuentas = cuentas.filter(tipo=tipo)
    hoy = timezone.now().date()
    saldos = contable.acumular_arbol(contable.movimientos_por_cuenta(None, hoy))
    filas = []
    for c in cuentas:
        d, h = saldos.get(c.pk, (contable.CERO, contable.CERO))
        filas.append({'c': c, 'saldo': contable.saldo_natural(c, d, h), 'tiene_mov': (d or h) != 0})
    return render(request, 'contabilidad/cuentas_lista.html', _ctx(
        request, filas=filas, filtros={'q': q, 'tipo': tipo}, tipos=TipoCuenta.choices,
        vacio=not CuentaContable.objects.exists(), total=len(filas)))


@login_required
def cuenta_form(request, pk=None):
    _gate(request.user, 'administrar_cuentas')
    cuenta = get_object_or_404(CuentaContable, pk=pk) if pk else None
    if request.method == 'POST':
        form = CuentaContableForm(request.POST, instance=cuenta)
        if form.is_valid():
            o = form.save(commit=False)
            if not o.pk:
                o.creado_por = request.user
            o.modificado_por = request.user
            o.save()
            messages.success(request, f'Cuenta {o.codigo} guardada.')
            return redirect('contabilidad:cuentas_lista')
    else:
        inicial = {}
        padre = request.GET.get('padre')
        if padre and not cuenta:
            p = CuentaContable.objects.filter(pk=padre).first()
            if p:
                n = p.hijas.count() + 1
                inicial = {'padre': p, 'tipo': p.tipo, 'naturaleza': p.naturaleza,
                           'codigo': f'{p.codigo}.{n:03d}' if p.nivel >= 3 else f'{p.codigo}.{n:02d}'}
        form = CuentaContableForm(instance=cuenta, initial=inicial)
    return render(request, 'contabilidad/cuenta_form.html', _ctx(request, form=form, cuenta=cuenta))


@login_required
@require_POST
def cuentas_cargar_base(request):
    _gate(request.user, 'administrar_cuentas')
    try:
        n = contable.cargar_catalogo_base(request.user)
        messages.success(request, f'Catálogo base cargado: {n} cuentas. Ajústalo a la realidad de ANA (agregar, renombrar, desactivar).')
    except ValidationError as e:
        messages.error(request, '; '.join(e.messages))
    return redirect('contabilidad:cuentas_lista')


@login_required
def cuentas_json(request):
    """Cuentas de detalle activas para el buscador del formulario de partidas."""
    _gate(request.user, 'ver_cuentas')
    data = [{'id': c.pk, 'codigo': c.codigo, 'nombre': c.nombre, 'tipo': c.get_tipo_display()}
            for c in CuentaContable.objects.filter(es_detalle=True, activo=True)]
    return JsonResponse({'cuentas': data})


# ─────────────────────────── Centros de costo ───────────────────────────
@login_required
def centros_lista(request):
    _gate(request.user, 'ver_cuentas')
    centros = CentroCosto.objects.annotate(n=Count('movimientos'))
    return render(request, 'contabilidad/centros_lista.html', _ctx(request, lista=centros))


@login_required
def centro_form(request, pk=None):
    _gate(request.user, 'administrar_cuentas')
    obj = get_object_or_404(CentroCosto, pk=pk) if pk else None
    form = CentroCostoForm(request.POST or None, instance=obj)
    if request.method == 'POST' and form.is_valid():
        o = form.save(commit=False)
        if not o.pk:
            o.creado_por = request.user
        o.modificado_por = request.user
        o.save()
        messages.success(request, f'Centro de costo {o.codigo} guardado.')
        return redirect('contabilidad:centros_lista')
    return render(request, 'contabilidad/centro_form.html', _ctx(request, form=form, obj=obj))


# ─────────────────────────── Períodos ───────────────────────────
@login_required
def periodos_lista(request):
    _gate(request.user, 'ver_partidas')
    hoy = timezone.now().date()
    contable.periodo_de(hoy)   # asegura que exista el mes actual
    periodos = (PeriodoContable.objects
                .annotate(n=Count('asientos', filter=Q(asientos__estado=EstadoAsiento.CONTABILIZADO)),
                          borradores=Count('asientos', filter=Q(asientos__estado=EstadoAsiento.BORRADOR))))
    return render(request, 'contabilidad/periodos_lista.html', _ctx(request, periodos=periodos,
                                                                     p_cerrar=puede(request.user, 'administrar_configuracion')))


@login_required
@require_POST
def periodo_toggle(request, pk):
    _gate(request.user, 'administrar_configuracion')
    p = get_object_or_404(PeriodoContable, pk=pk)
    try:
        if p.abierto:
            contable.cerrar_periodo(p, request.user); messages.success(request, f'Período {p} cerrado.')
        else:
            contable.reabrir_periodo(p, request.user); messages.warning(request, f'Período {p} reabierto.')
    except ValidationError as e:
        messages.error(request, '; '.join(e.messages))
    return redirect('contabilidad:periodos_lista')


# ─────────────────────────── Partidas ───────────────────────────
@login_required
def asientos_lista(request):
    _gate(request.user, 'ver_partidas')
    desde, hasta = _rango(request)
    estado = request.GET.get('estado') or ''
    tipo = request.GET.get('tipo') or ''
    q = (request.GET.get('q') or '').strip()
    qs = Asiento.objects.filter(fecha__gte=desde, fecha__lte=hasta)
    if estado:
        qs = qs.filter(estado=estado)
    if tipo:
        qs = qs.filter(tipo=tipo)
    if q:
        qs = qs.filter(Q(numero__icontains=q) | Q(descripcion__icontains=q) | Q(referencia__icontains=q))
    qs = qs.select_related('periodo', 'creado_por')
    tot = qs.filter(estado=EstadoAsiento.CONTABILIZADO).aggregate(d=Sum('total_debe'))['d'] or contable.CERO
    return render(request, 'contabilidad/asientos_lista.html', _ctx(
        request, asientos=qs[:500], n=qs.count(), total=tot,
        filtros={'desde': desde, 'hasta': hasta, 'estado': estado, 'tipo': tipo, 'q': q},
        estados=EstadoAsiento.choices, tipos=TipoAsiento.choices,
        borradores=Asiento.objects.filter(estado=EstadoAsiento.BORRADOR).count()))


def _lineas_post(request):
    """Líneas del formulario: arrays paralelos cuenta[], centro[], desc[], debe[], haber[]."""
    ids = request.POST.getlist('cuenta[]')
    centros = request.POST.getlist('centro[]')
    descs = request.POST.getlist('desc[]')
    debes = request.POST.getlist('debe[]')
    habers = request.POST.getlist('haber[]')
    cuentas = {c.pk: c for c in CuentaContable.objects.filter(pk__in=[i for i in ids if i.isdigit()])}
    cc = {c.pk: c for c in CentroCosto.objects.all()}
    lineas = []
    for i, cid in enumerate(ids):
        lineas.append({'cuenta': cuentas.get(int(cid)) if cid.isdigit() else None,
                       'centro_costo': cc.get(int(centros[i])) if i < len(centros) and centros[i].isdigit() else None,
                       'descripcion': descs[i] if i < len(descs) else '',
                       'debe': debes[i] if i < len(debes) else '0', 'haber': habers[i] if i < len(habers) else '0'})
    return lineas


@login_required
def asiento_form(request, pk=None):
    asiento = get_object_or_404(Asiento, pk=pk) if pk else None
    _gate(request.user, 'modificar_partidas' if asiento else 'crear_partidas')
    if asiento and not asiento.editable:
        messages.error(request, 'Esta partida ya no está en borrador; no se puede editar.')
        return redirect('contabilidad:asiento_detalle', pk=asiento.pk)
    lineas_ini = []
    if request.method == 'POST':
        form = AsientoHeaderForm(request.POST, instance=asiento)
        lineas = _lineas_post(request)
        lineas_ini = [{'cuenta_id': l['cuenta'].pk if l['cuenta'] else '', 'codigo': l['cuenta'].codigo if l['cuenta'] else '',
                       'nombre': l['cuenta'].nombre if l['cuenta'] else '', 'centro': l['centro_costo'].pk if l['centro_costo'] else '',
                       'desc': l['descripcion'], 'debe': l['debe'], 'haber': l['haber']} for l in lineas]
        if form.is_valid():
            o = form.save(commit=False)
            try:
                contable.guardar_asiento(o, lineas, request.user, contabilizar=request.POST.get('accion') == 'contabilizar')
                messages.success(request, f'Partida {o.numero or "(borrador)"} {"contabilizada" if o.estado == EstadoAsiento.CONTABILIZADO else "guardada como borrador"}.')
                return redirect('contabilidad:asiento_detalle', pk=o.pk)
            except ValidationError as e:
                messages.error(request, '; '.join(e.messages))
    else:
        form = AsientoHeaderForm(instance=asiento, initial=None if asiento else {'fecha': timezone.now().date()})
        if asiento:
            lineas_ini = [{'cuenta_id': d.cuenta_id, 'codigo': d.cuenta.codigo, 'nombre': d.cuenta.nombre,
                           'centro': d.centro_costo_id or '', 'desc': d.descripcion, 'debe': str(d.debe), 'haber': str(d.haber)}
                          for d in asiento.detalles.select_related('cuenta')]
    return render(request, 'contabilidad/asiento_form.html', _ctx(request, form=form, asiento=asiento, lineas_ini=lineas_ini))


@login_required
def asiento_detalle(request, pk):
    _gate(request.user, 'ver_partidas')
    a = get_object_or_404(Asiento.objects.select_related('periodo', 'creado_por', 'contabilizado_por', 'anulado_por'), pk=pk)
    return render(request, 'contabilidad/asiento_detalle.html', _ctx(
        request, a=a, detalles=a.detalles.select_related('cuenta', 'centro_costo'),
        p_anular=puede(request.user, 'modificar_partidas'), p_periodo_abierto=a.periodo.abierto))


@login_required
@require_POST
def asiento_contabilizar(request, pk):
    _gate(request.user, 'crear_partidas')
    a = get_object_or_404(Asiento, pk=pk)
    try:
        contable.contabilizar_asiento(a, request.user)
        messages.success(request, f'Partida {a.numero} contabilizada.')
    except ValidationError as e:
        messages.error(request, '; '.join(e.messages))
    return redirect('contabilidad:asiento_detalle', pk=pk)


@login_required
@require_POST
def asiento_anular(request, pk):
    _gate(request.user, 'modificar_partidas')
    a = get_object_or_404(Asiento, pk=pk)
    try:
        contable.anular_asiento(a, request.user, request.POST.get('motivo', ''))
        messages.warning(request, f'Partida {a.numero} anulada.')
    except ValidationError as e:
        messages.error(request, '; '.join(e.messages))
    return redirect('contabilidad:asiento_detalle', pk=pk)


@login_required
@require_POST
def asiento_eliminar(request, pk):
    """Solo borradores (nunca se borra una partida contabilizada o anulada)."""
    _gate(request.user, 'modificar_partidas')
    a = get_object_or_404(Asiento, pk=pk)
    if not a.editable:
        messages.error(request, 'Solo se eliminan borradores.')
        return redirect('contabilidad:asiento_detalle', pk=pk)
    a.delete()
    messages.success(request, 'Borrador eliminado.')
    return redirect('contabilidad:asientos_lista')


# ─────────────────────────── Mayor y reportes ───────────────────────────
@login_required
def mayor(request):
    _gate(request.user, 'ver_cuentas')
    desde, hasta = _rango(request)
    centro = _centro(request)
    cid = request.GET.get('cuenta')
    cuenta = CuentaContable.objects.filter(pk=cid).first() if cid else None
    data = contable.mayor_cuenta(cuenta, desde, hasta, centro) if cuenta else None
    return render(request, 'contabilidad/mayor.html', _ctx(
        request, cuenta=cuenta, data=data, cuentas=CuentaContable.objects.filter(activo=True),
        filtros={'desde': desde, 'hasta': hasta, 'cuenta': cid or '', 'centro': centro.pk if centro else ''}))


@login_required
def reportes(request):
    _gate(request.user, 'ver_reportes')
    return render(request, 'contabilidad/reportes.html', _ctx(request))


@login_required
def balanza(request):
    _gate(request.user, 'ver_reportes')
    desde, hasta = _rango(request); centro = _centro(request)
    solo_detalle = request.GET.get('detalle') == '1'
    filas, tot = contable.balanza_comprobacion(desde, hasta, centro, solo_detalle)
    if request.GET.get('xlsx'):
        cols = ['Código', 'Cuenta', 'Saldo inicial Debe', 'Saldo inicial Haber', 'Debe', 'Haber', 'Saldo final Debe', 'Saldo final Haber']
        rows = [[f['cuenta'].codigo, f['cuenta'].nombre, f['ini_d'], f['ini_h'], f['debe'], f['haber'], f['fin_d'], f['fin_h']] for f in filas]
        rows.append(['', 'TOTALES', tot['ini_d'], tot['ini_h'], tot['debe'], tot['haber'], tot['fin_d'], tot['fin_h']])
        return _xlsx(f'Balanza de comprobación {desde:%d/%m/%Y} – {hasta:%d/%m/%Y}', cols, rows, 'balanza_comprobacion')
    return render(request, 'contabilidad/balanza.html', _ctx(
        request, filas=filas, tot=tot, filtros={'desde': desde, 'hasta': hasta, 'centro': centro.pk if centro else '', 'detalle': solo_detalle}))


@login_required
def balance_general(request):
    _gate(request.user, 'ver_reportes')
    hasta = _fecha(request.GET.get('hasta'), timezone.now().date()); centro = _centro(request)
    data = contable.balance_general(hasta, centro)
    if request.GET.get('xlsx'):
        rows = []
        for sec, key in (('ACTIVO', 'activo'), ('PASIVO', 'pasivo'), ('PATRIMONIO', 'patrimonio')):
            rows.append([sec, '', ''])
            rows += [[f['cuenta'].codigo, f['cuenta'].nombre, f['saldo']] for f in data[key]]
        rows.append(['', 'Resultado del ejercicio', data['resultado']])
        rows.append(['', 'TOTAL ACTIVO', data['total_activo']]); rows.append(['', 'TOTAL PASIVO + PATRIMONIO', data['total_pasivo_patrimonio']])
        return _xlsx(f'Balance general al {hasta:%d/%m/%Y}', ['Código', 'Cuenta', 'Saldo'], rows, 'balance_general')
    return render(request, 'contabilidad/balance_general.html', _ctx(request, d=data, hasta=hasta, centro=centro))


@login_required
def estado_resultados(request):
    _gate(request.user, 'ver_reportes')
    desde, hasta = _rango(request); centro = _centro(request)
    data = contable.estado_resultados(desde, hasta, centro)
    if request.GET.get('xlsx'):
        rows = []
        for sec, key in (('INGRESOS', 'ingresos'), ('COSTOS', 'costos'), ('GASTOS', 'gastos')):
            rows.append([sec, '', ''])
            rows += [[f['cuenta'].codigo, f['cuenta'].nombre, f['saldo']] for f in data[key]]
        rows += [['', 'Utilidad bruta', data['utilidad_bruta']], ['', 'RESULTADO DEL PERÍODO', data['resultado']]]
        return _xlsx(f'Estado de resultados {desde:%d/%m/%Y} – {hasta:%d/%m/%Y}', ['Código', 'Cuenta', 'Importe'], rows, 'estado_resultados')
    return render(request, 'contabilidad/estado_resultados.html', _ctx(request, d=data, desde=desde, hasta=hasta, centro=centro))


@login_required
def libro_diario(request):
    _gate(request.user, 'ver_partidas')
    desde, hasta = _rango(request)
    asientos = (Asiento.objects.filter(estado=EstadoAsiento.CONTABILIZADO, fecha__gte=desde, fecha__lte=hasta)
                .prefetch_related('detalles__cuenta', 'detalles__centro_costo').order_by('fecha', 'numero'))
    if request.GET.get('xlsx'):
        rows = []
        for a in asientos:
            for d in a.detalles.all():
                rows.append([a.numero, a.fecha, a.get_tipo_display(), a.descripcion, d.cuenta.codigo, d.cuenta.nombre,
                             d.centro_costo.codigo if d.centro_costo else '', d.descripcion, d.debe, d.haber])
        return _xlsx(f'Libro diario {desde:%d/%m/%Y} – {hasta:%d/%m/%Y}',
                     ['Partida', 'Fecha', 'Tipo', 'Concepto', 'Código', 'Cuenta', 'Centro', 'Detalle', 'Debe', 'Haber'], rows, 'libro_diario')
    tot = asientos.aggregate(d=Sum('total_debe'))['d'] or contable.CERO
    return render(request, 'contabilidad/libro_diario.html', _ctx(request, asientos=asientos, total=tot, filtros={'desde': desde, 'hasta': hasta}))


def _xlsx(titulo, cols, rows, archivo):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = archivo[:31]
    ws.append([titulo]); ws['A1'].font = Font(bold=True, size=13)
    ws.append(['Asociación Nuevo Amanecer · Sistema TechCare · ' + timezone.now().strftime('%d/%m/%Y %H:%M')])
    ws.append([]); ws.append(cols)
    for c in ws[4]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
    for r in rows:
        ws.append(r)
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14 if i > 2 else (14 if i == 1 else 46)
    for row in ws.iter_rows(min_row=5):
        for c in row:
            if isinstance(c.value, (int, float)) or hasattr(c.value, 'quantize'):
                c.number_format = '#,##0.00'; c.alignment = Alignment(horizontal='right')
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    resp = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{archivo}.xlsx"'
    return resp
