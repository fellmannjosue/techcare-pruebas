# <--- hecho por claude code (01-oct-2026): Mantenimiento Mecánico — control de combustibles y aceites.
# Una página por insumo (gasolina, diésel, aceite 40/50/hidráulico): saldo, entradas, entregas (formularios),
# maquinaria agregable, PDF de ficha, reporte tabla y reporte general en PDF y Excel.
import datetime
import json
import os
import re
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.staticfiles import finders
from django.core.exceptions import PermissionDenied
from django.db.models import Sum, Count
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.views.decorators.http import require_POST

from . import services
from .models import Comentario, Entrada, Entrega, FotoLimpiezaTanque, Insumo, LecturaBarril, LimpiezaTanque, Maquinaria, FIRMA_AUDITORA, GRUPO_ACCESO, GRUPO_LECTURA, LITROS_POR_GALON

MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio',
         'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']


# ── Acceso ──────────────────────────────────────────────────────────────────
def puede_editar(user):
    """Registra entregas/entradas, configura y agrega maquinaria."""
    return user.is_superuser or user.groups.filter(name=GRUPO_ACCESO).exists()


def puede_modificar(user):
    """<--- hecho por claude code (02-oct-2026): editar/eliminar entregas y entradas SOLO el superusuario (lo registrado
    no se cambia: es lo que audita la auditora). El grupo mantenimiento_mecanico solo REGISTRA."""
    return user.is_superuser


def puede_ver(user):
    """<--- hecho por claude code (02-oct-2026): el grupo de lectura (auditoría) ve todo y comenta, pero no modifica."""
    return puede_editar(user) or user.groups.filter(name=GRUPO_LECTURA).exists()


def _gate(request, escribir=False, modificar=False):
    ok = puede_modificar(request.user) if modificar else (puede_editar(request.user) if escribir else puede_ver(request.user))
    if not ok:
        raise PermissionDenied('No tienes permiso para esta acción en Mantenimiento Mecánico.')


def _insumo(clave):
    return get_object_or_404(Insumo, clave=clave, activo=True)


def _dec(v, default=None):
    try:
        return Decimal(str(v).replace(',', '.').strip())
    except (InvalidOperation, ValueError, AttributeError):
        return default


def _fecha(v, default=None):
    try:
        return datetime.date.fromisoformat((v or '').strip())
    except ValueError:
        return default


def _hora(v):
    v = (v or '').strip()
    for fmt in ('%H:%M', '%H:%M:%S'):
        try:
            return datetime.datetime.strptime(v, fmt).time()
        except ValueError:
            pass
    return None


def _rango(request):
    """?desde=&hasta= (ISO). Por defecto el mes actual."""
    hoy = datetime.date.today()
    desde = _fecha(request.GET.get('desde'), hoy.replace(day=1))
    hasta = _fecha(request.GET.get('hasta'), hoy)
    if hasta < desde:
        desde, hasta = hasta, desde
    return desde, hasta


def _ctx_base(request, insumo=None):
    insumos = list(Insumo.objects.filter(activo=True))
    return {
        'insumos': insumos,
        'insumo': insumo,
        'combustibles': [i for i in insumos if i.categoria == 'combustible'],
        'aceites': [i for i in insumos if i.categoria == 'aceite'],
    }


# ── Páginas ─────────────────────────────────────────────────────────────────
@login_required
def index(request):
    """<--- hecho por claude code (12-oct-2026): SUB-DASHBOARD del módulo (pedido del usuario: combustibles/aceites y
    vehículos unidos). Cada sección sale según el permiso: insumos (grupos mantenimiento_mecanico*) y vehículos (vehiculos_*)."""
    from . import views_vehiculos as VV, vehiculos_db as VS
    ver_insumos, ver_veh = puede_ver(request.user), VV.puede_ver(request.user)
    if not (ver_insumos or ver_veh):
        raise PermissionDenied('No tienes permiso para Mantenimiento Mecánico.')
    ctx = {'ver_insumos': ver_insumos, 'ver_veh': ver_veh}
    if ver_insumos:
        hoy = datetime.date.today()
        mes = {r['insumo']: r['t'] for r in Entrega.objects.filter(fecha__year=hoy.year, fecha__month=hoy.month)
               .values('insumo').annotate(t=Sum('galones'))}
        tiles = []
        for i in Insumo.objects.filter(activo=True):
            r = services.resumen(i)
            tiles.append({'i': i, 'r': r, 'mes': mes.get(i.pk) or Decimal('0')})
        ctx.update({'tiles_comb': [t for t in tiles if t['i'].categoria == 'combustible'],
                    'tiles_aceite': [t for t in tiles if t['i'].categoria != 'combustible'],
                    'n_bajos': sum(1 for t in tiles if t['r']['bajo']), 'hoy': hoy})
    if ver_veh:
        vehs = VS.vehiculos()
        orden = {'vencido': 0, 'pronto': 1, 'ok': 2, 'sin_dato': 3}
        vehs.sort(key=lambda v: (orden[v['semaforo']], v['falta'] if v['falta'] is not None else 10**9))
        ctx.update({'veh_total': len(vehs), 'veh_vencidos': sum(1 for v in vehs if v['semaforo'] == 'vencido'),
                    'veh_pronto': sum(1 for v in vehs if v['semaforo'] == 'pronto'),
                    'veh_ok': sum(1 for v in vehs if v['semaforo'] == 'ok'),
                    'veh_atencion': [v for v in vehs if v['semaforo'] in ('vencido', 'pronto')][:8],
                    'puede_registrar_veh': VV.puede_registrar(request.user)})
    return render(request, 'mantenimiento_mecanico/hub.html', ctx)


@login_required
def pagina(request, clave):
    _gate(request)
    insumo = _insumo(clave)
    desde, hasta = _rango(request)
    r = services.resumen(insumo)
    entregas = (Entrega.objects.filter(insumo=insumo, fecha__range=(desde, hasta))
                .select_related('maquinaria', 'vehiculo', 'registrado_por').prefetch_related('comentarios__autor')
                .order_by('-fecha', '-hora', '-id'))
    entradas = Entrada.objects.filter(insumo=insumo).select_related('registrado_por').order_by('-fecha', '-id')[:30]
    hoy = datetime.date.today()
    mes_qs = Entrega.objects.filter(insumo=insumo, fecha__year=hoy.year, fecha__month=hoy.month)
    maquinaria = list(Maquinaria.objects.filter(activo=True).values('id', 'nombre', 'usa_kilometraje'))
    ctx = _ctx_base(request, insumo)
    ctx.update({
        'resumen': r,
        'entregas': entregas,
        'entradas': entradas,
        'lecturas': _lecturas_barril(insumo, r['saldo']),
        # <--- hecho por claude code (19-oct-2026): mantenimiento del tanque (limpieza del tanque y del filtro, cada 4 meses)
        'limpieza': services.estado_limpieza(insumo),
        'limpiezas': (LimpiezaTanque.objects.filter(insumo=insumo).select_related('registrado_por').prefetch_related('fotos')[:20]
                      if insumo.presentacion == 'tanque' else []),   # <--- hecho por claude code (14-oct-2026): lectura aproximada (solo referencia)
        'niveles_barril': [(k, l, (LecturaBarril.FRACCION[k] * (insumo.capacidad_envase or Decimal('55'))).quantize(Decimal('0.1')))
                           for k, l in LecturaBarril.NIVELES] if insumo.presentacion == 'barril' else [],
        'desde': desde, 'hasta': hasta,
        'kpi_mes_n': mes_qs.count(),
        'kpi_mes_gal': mes_qs.aggregate(t=Sum('galones'))['t'] or Decimal('0'),
        'maquinaria': maquinaria,
        'vehiculos_diesel': _vehiculos_diesel() if insumo.clave == 'diesel' else [],   # <--- hecho por claude code (12-oct-2026)
        # <--- hecho por claude code (01-oct-2026): listas tal cual; json_script ya serializa (con json.dumps salían como TEXTO y el JS se caía)
        'maquinaria_json': maquinaria,
        'entregas_json': [{
            'id': e.pk, 'fecha': e.fecha.isoformat(), 'hora': e.hora.strftime('%H:%M') if e.hora else '',
            'maquinaria': e.maquinaria_id or '', 'vehiculo': e.vehiculo_id or '', 'kilometraje': e.kilometraje or '',
            'galones': str(e.galones), 'persona': e.persona, 'observaciones': e.observaciones, 'codigo': e.codigo,
            'comentarios': [{'autor': (c.autor.get_full_name() or c.autor.username) if c.autor else '—',
                             'fecha': c.creado.strftime('%d/%m/%Y %H:%M'), 'texto': c.texto} for c in e.comentarios.all()],
        } for e in entregas],
        'puede_editar': puede_editar(request.user),
        'puede_modificar': puede_modificar(request.user),
        'firma_auditora': FIRMA_AUDITORA[0],
        'responsable_default': request.user.get_full_name() or request.user.username,
        'hoy': hoy,
    })
    return render(request, 'mantenimiento_mecanico/pagina.html', ctx)


# ── Entregas (formularios) ───────────────────────────────────────────────────
def _vehiculos_diesel():
    """Vehículos en uso para cargar diésel (con su último km, para la sugerencia del formulario)."""
    from . import vehiculos_db
    return [{'id': v['id'], 'nombre': v['nombre'], 'placa': v['Placa'], 'km': v['estado_actual']}
            for v in vehiculos_db.vehiculos()]


def _por_equipo(qs):
    """Consumo agrupado por equipo: vehículo (diésel) o maquinaria. Misma clave que antes para no tocar las plantillas."""
    tot = {}
    for e in qs.select_related('vehiculo', 'maquinaria'):
        k = e.equipo_nombre or ''
        t = tot.setdefault(k, {'maquinaria__nombre': k, 'n': 0, 'gal': Decimal('0')})
        t['n'] += 1
        t['gal'] += e.galones
    return sorted(tot.values(), key=lambda x: -x['gal'])


def _leer_entrega(request, insumo, obj=None):
    """Valida el POST del formulario de entrega. Devuelve (obj, errores)."""
    err = []
    fecha = _fecha(request.POST.get('fecha'))
    if not fecha:
        err.append('La fecha es obligatoria.')
    gal = _dec(request.POST.get('galones'))
    if gal is None or gal <= 0:
        err.append('Los galones deben ser mayores a cero.')
    maq = veh = None
    km = (request.POST.get('kilometraje') or '').strip()
    km_val = int(km) if km.isdigit() else None
    if insumo.clave == 'diesel':
        # <--- hecho por claude code (12-oct-2026): el diésel se carga a un vehículo (como en Access: Gasto · EstadoKm · Fecha · Comentario)
        from .models import VehiculoMant
        vid = (request.POST.get('vehiculo') or '').strip()
        veh = VehiculoMant.objects.filter(pk=int(vid)).first() if vid.isdigit() else None
        if not veh:
            err.append('Elige el vehículo al que se carga el diésel.')
        elif km_val is not None and obj is None and not veh.sin_tacometro:
            from . import vehiculos_db
            est = vehiculos_db.vehiculo(veh.pk)['estado_actual']
            if est is not None and km_val < est:
                err.append(f'El kilometraje ({km_val:,}) es menor al último registrado de {veh.nombre} ({est:,}).'.replace(',', '.'))
    else:
        mid = (request.POST.get('maquinaria') or '').strip()
        if mid.isdigit():
            maq = Maquinaria.objects.filter(pk=int(mid)).first()
    if insumo.requiere_kilometraje and km_val is None:
        err.append('El kilometraje es obligatorio en diésel.')
    if km_val is not None and maq:
        ant = services.kilometraje_anterior(maq, excluir_pk=obj.pk if obj else None)
        if ant is not None and km_val < ant:
            err.append(f'El kilometraje ({km_val}) es menor al último registrado de {maq.nombre} ({ant}).')
    if err:
        return None, err
    obj = obj or Entrega(insumo=insumo, registrado_por=request.user)
    obj.fecha = fecha
    obj.hora = _hora(request.POST.get('hora'))
    obj.maquinaria = maq
    obj.vehiculo = veh
    obj.kilometraje = km_val
    obj.galones = gal
    obj.persona = (request.POST.get('persona') or '').strip()[:120]
    obj.observaciones = (request.POST.get('observaciones') or '').strip()
    obj.firma_responsable = ((request.POST.get('firma_responsable') or '').strip()
                             or request.user.get_full_name() or request.user.username)[:120]
    obj.firma_auditora = FIRMA_AUDITORA[0]
    return obj, []


def _volver(insumo, request):
    url = reverse('mantenimiento_mecanico:pagina', args=[insumo.clave])
    qs = request.POST.get('_qs') or ''
    return redirect(url + ('?' + qs if qs else ''))


@login_required
@require_POST
def entrega_crear(request, clave):
    _gate(request, escribir=True)
    insumo = _insumo(clave)
    obj, err = _leer_entrega(request, insumo)
    if err:
        for e in err:
            messages.error(request, e)
        return _volver(insumo, request)
    antes = services.saldo(insumo)
    obj.numero = services.siguiente_numero(insumo)
    obj.save()
    despues = antes - obj.galones
    n = services.alertar_si_bajo(insumo, antes, despues, url=reverse('mantenimiento_mecanico:pagina', args=[insumo.clave]))
    msg = f'Entrega {obj.codigo} registrada: {obj.galones} gal. Saldo: {despues:.2f} gal.'
    if despues < 0:
        msg += ' ⚠ El saldo quedó negativo: registra la entrada pendiente.'
    messages.success(request, msg)
    if n:
        messages.warning(request, f'Saldo bajo de {insumo.nombre}: se avisó a {n} persona(s).')
    # <--- hecho por claude code (12-oct-2026): el km de la carga de diésel queda como LECTURA del vehículo (alimenta el
    # conteo del próximo mantenimiento y, si entra en "pronto"/"vencido", avisa al grupo vehiculos_alertas)
    if obj.vehiculo_id and obj.kilometraje is not None:
        from . import vehiculos_db
        from .models import LecturaVehiculo
        from .views_vehiculos import _avisar
        v_antes = vehiculos_db.vehiculo(obj.vehiculo_id)
        LecturaVehiculo.objects.create(vehiculo_id=obj.vehiculo_id, valor=obj.kilometraje, fecha=obj.fecha,
                                       registrado_por=request.user, usuario=(request.user.first_name or request.user.username)[:50])
        n2 = _avisar(v_antes, vehiculos_db.vehiculo(obj.vehiculo_id), request.user.get_full_name() or request.user.username)
        if n2:
            messages.warning(request, f'A {obj.vehiculo.nombre} le toca mantenimiento: se avisó a {n2} persona(s).')
    return _volver(insumo, request)


@login_required
@require_POST
def entrega_editar(request, pk):
    _gate(request, modificar=True)
    obj = get_object_or_404(Entrega, pk=pk)
    insumo = obj.insumo
    obj2, err = _leer_entrega(request, insumo, obj)
    if err:
        for e in err:
            messages.error(request, e)
        return _volver(insumo, request)
    obj2.save()
    messages.success(request, f'Entrega {obj2.codigo} actualizada.')
    return _volver(insumo, request)




@login_required
@require_POST
def comentario_crear(request, pk):
    """<--- hecho por claude code (02-oct-2026): comentario de auditoría (lo puede dejar quien VE el módulo)."""
    _gate(request)
    e = get_object_or_404(Entrega, pk=pk)
    texto = (request.POST.get('texto') or '').strip()
    if len(texto) < 2:
        messages.error(request, 'Escribe el comentario.')
    else:
        Comentario.objects.create(entrega=e, autor=request.user, texto=texto[:2000])
        messages.success(request, f'Comentario agregado a {e.codigo}.')
    return _volver(e.insumo, request)


# ── Entradas (inventario) ────────────────────────────────────────────────────
@login_required
@require_POST
def entrada_crear(request, clave):
    _gate(request, escribir=True)
    insumo = _insumo(clave)
    fecha = _fecha(request.POST.get('fecha'))
    cant = _dec(request.POST.get('cantidad'))
    # <--- hecho por claude code (14-oct-2026): gasolina y diésel se compran en LITROS → los galones se calculan aquí
    litros = _dec(request.POST.get('litros')) if insumo.categoria == 'combustible' else None
    # <--- hecho por claude code (21-oct-2026): compra extra en galones (combustible). cantidad = galones de litros + extra.
    extra = _dec(request.POST.get('galones_extra')) if insumo.categoria == 'combustible' else None
    if extra is not None and extra <= 0:
        extra = None
    if litros is not None and litros > 0:
        cant = (litros / LITROS_POR_GALON).quantize(Decimal('0.01')) + (extra or Decimal('0'))
    else:
        litros = None
        if extra is not None:   # solo compra extra, sin litros: el total es el extra
            cant = extra
    tipo = request.POST.get('tipo') or {'galones': 'compra', 'tanque': 'llenado', 'barril': 'barril'}[insumo.presentacion]
    if tipo not in dict(Entrada.TIPOS):
        tipo = 'ajuste'
    if not fecha or cant is None or (cant <= 0 and tipo != 'ajuste'):
        messages.error(request, 'Fecha y cantidad (mayor a cero) son obligatorias.')
        return _volver(insumo, request)
    if insumo.presentacion == 'tanque' and insumo.capacidad_envase:
        nuevo = services.saldo(insumo) + cant
        if nuevo > insumo.capacidad_envase and request.POST.get('confirmar_capacidad') != '1':
            messages.error(request, f'Con este llenado el tanque quedaría en {nuevo:.2f} gal, más que su capacidad '
                                    f'({insumo.capacidad_envase} gal). Marca "Confirmar de todas formas" si es correcto.')
            return _volver(insumo, request)
    Entrada.objects.create(insumo=insumo, fecha=fecha, cantidad=cant, litros=litros, galones_extra=extra, tipo=tipo,
                           observaciones=(request.POST.get('observaciones') or '').strip(),
                           registrado_por=request.user)
    messages.success(request, f'{dict(Entrada.TIPOS)[tipo]}: +{cant} gal' + (f' ({litros} L' + (f' + {extra} gal extra' if extra else '') + ')' if litros else (f' ({extra} gal extra)' if extra else '')) +
                     f'. Saldo: {services.saldo(insumo):.2f} gal.')
    return _volver(insumo, request)




def _lecturas_barril(insumo, saldo):
    """Últimas lecturas aproximadas del barril; la primera trae la diferencia contra el saldo calculado de HOY."""
    if insumo.presentacion != 'barril':
        return []
    lecturas = list(LecturaBarril.objects.filter(insumo=insumo).select_related('registrado_por')[:10])
    if lecturas:
        lecturas[0].diferencia = lecturas[0].galones_aprox - saldo
    return lecturas


@login_required
@require_POST
def lectura_crear(request, clave):
    """<--- hecho por claude code (14-oct-2026): lectura APROXIMADA del barril (aceites y coolant). SOLO REFERENCIA: el saldo no cambia."""
    _gate(request, escribir=True)
    insumo = _insumo(clave)
    if insumo.presentacion != 'barril':
        raise PermissionDenied
    fecha = _fecha(request.POST.get('fecha'))
    nivel = request.POST.get('nivel')
    try:
        llenos = max(0, min(int(request.POST.get('barriles_llenos') or 0), 99))
    except ValueError:
        llenos = 0
    if not fecha or nivel not in LecturaBarril.FRACCION:
        messages.error(request, 'Elige el nivel del barril y la fecha.')
        return _volver(insumo, request)
    cap = insumo.capacidad_envase or Decimal('55')
    gal = ((llenos + LecturaBarril.FRACCION[nivel]) * cap).quantize(Decimal('0.01'))
    lec = LecturaBarril.objects.create(insumo=insumo, fecha=fecha, nivel=nivel, barriles_llenos=llenos, galones_aprox=gal,
                                       observaciones=(request.POST.get('observaciones') or '').strip()[:255],
                                       registrado_por=request.user)
    messages.success(request, f'Lectura aproximada registrada: {lec.get_nivel_display()}'
                              + (f' + {llenos} barril(es) lleno(s)' if llenos else '')
                              + f' ≈ {gal} gal. El saldo calculado no cambia ({services.saldo(insumo):.2f} gal).')
    return _volver(insumo, request)


@login_required
@require_POST
def limpieza_crear(request, clave):
    """<--- hecho por claude code (19-oct-2026): registra el MANTENIMIENTO DEL TANQUE. El saldo NO cambia (el diésel se
    saca y se vuelve a meter); la próxima limpieza pasa a ser esta fecha + 4 meses."""
    from django.db import transaction
    _gate(request, escribir=True)
    insumo = _insumo(clave)
    if insumo.presentacion != 'tanque':
        raise PermissionDenied
    fecha = _fecha(request.POST.get('fecha'))
    responsable = (request.POST.get('responsable') or '').strip()[:150]
    marcas = {k: request.POST.get(k) == '1' for k in ('tanque_limpiado', 'filtro_limpiado', 'filtro_cambiado')}
    if not fecha or not responsable or not any(marcas.values()):
        messages.error(request, 'Indica la fecha, quién lo hizo y al menos un trabajo (tanque o filtro).')
        return _volver(insumo, request)
    if fecha > datetime.date.today():
        messages.error(request, 'La fecha de la limpieza no puede ser futura.')
        return _volver(insumo, request)
    fotos = [f for f in request.FILES.getlist('fotos') if (f.content_type or '').startswith('image/')][:8]
    with transaction.atomic():
        lim = LimpiezaTanque.objects.create(
            insumo=insumo, codigo=services.siguiente_codigo_limpieza(), fecha=fecha, responsable=responsable,
            galones_sacados=_dec(request.POST.get('galones_sacados')),
            galones_reingresados=_dec(request.POST.get('galones_reingresados')),
            observaciones=(request.POST.get('observaciones') or '').strip()[:2000], registrado_por=request.user, **marcas)
        for f in fotos:
            FotoLimpiezaTanque.objects.create(limpieza=lim, imagen=f)
    insumo.aviso_limpieza = ''
    insumo.save(update_fields=['aviso_limpieza'])
    e = services.estado_limpieza(insumo)
    messages.success(request, f'Limpieza {lim.codigo} registrada. Próxima limpieza del tanque: {e["proxima"]:%d/%m/%Y} ({e["texto"]}). '
                              f'El saldo de {insumo.nombre} no cambia.')
    return _volver(insumo, request)


@login_required
def limpieza_pdf(request, pk):
    _gate(request)
    lim = get_object_or_404(LimpiezaTanque.objects.select_related('insumo', 'registrado_por').prefetch_related('fotos'), pk=pk)
    fotos = [f'file://{f.imagen.path}' for f in lim.fotos.all() if f.imagen and os.path.exists(f.imagen.path)]
    ctx = {'l': lim, 'insumo': lim.insumo, 'logo': _logo_file(), 'auditora': FIRMA_AUDITORA, 'fotos': fotos,
           'proxima': services.sumar_meses(lim.fecha, services.LIMPIEZA_MESES)}
    return _pdf(request, 'mantenimiento_mecanico/pdf/limpieza.html', ctx, f'{lim.codigo}.pdf')


@login_required
@require_POST
def insumo_config(request, clave):
    """Capacidad del envase y mínimo de alerta (desde el modal Configurar)."""
    _gate(request, escribir=True)
    insumo = _insumo(clave)
    cap = _dec(request.POST.get('capacidad_envase'))
    minimo = _dec(request.POST.get('minimo_alerta'))
    if insumo.presentacion != 'galones':
        insumo.capacidad_envase = cap if (cap is not None and cap > 0) else insumo.capacidad_envase
    if minimo is not None and minimo >= 0:
        insumo.minimo_alerta = minimo
    insumo.save()
    messages.success(request, 'Configuración guardada.')
    return _volver(insumo, request)


# ── Maquinaria (select agregable) ────────────────────────────────────────────
@login_required
@require_POST
def maquinaria_add(request):
    _gate(request, escribir=True)
    try:
        body = json.loads(request.body or b'{}')
    except ValueError:
        body = {}
    nombre = (body.get('nombre') or '').strip()[:120]
    if len(nombre) < 2:
        return JsonResponse({'ok': False, 'error': 'Escribe el nombre de la maquinaria.'}, status=400)
    m = Maquinaria.objects.filter(nombre__iexact=nombre).first()
    if m is None:
        m = Maquinaria.objects.create(nombre=nombre, tipo=body.get('tipo') if body.get('tipo') in dict(Maquinaria.TIPOS) else 'maquina',
                                      usa_kilometraje=bool(body.get('usa_kilometraje')), creado_por=request.user)
    elif not m.activo:
        m.activo = True
        m.save(update_fields=['activo'])
    return JsonResponse({'ok': True, 'id': m.pk, 'nombre': m.nombre, 'usa_kilometraje': m.usa_kilometraje})


# ── PDF ──────────────────────────────────────────────────────────────────────
def _pdf(request, template, ctx, filename, inline=True):
    """WeasyPrint con el CSS cargado DEL DISCO (ver AGENTS.md: por HTTP no llega desde el propio servidor)."""
    from weasyprint import HTML, CSS
    html = render_to_string(template, ctx, request=request)
    hojas = []
    for href in re.findall(r'<link[^>]+href="([^"]+\.css[^"]*)"', html):
        ruta = finders.find(href.split('?')[0].replace(settings.STATIC_URL, '', 1))
        if ruta:
            hojas.append(CSS(filename=ruta))
    pdf = HTML(string=html, base_url=request.build_absolute_uri('/')).write_pdf(stylesheets=hojas)
    resp = HttpResponse(pdf, content_type='application/pdf')
    resp['Content-Disposition'] = f'{"inline" if inline else "attachment"}; filename="{filename}"'
    return resp


def _logo_file():
    p = finders.find('mantenimiento_mecanico/logo_ana.png')   # <--- versión de 256 px: el original (2048 px, 1.8 MB) inflaba cada PDF
    return f'file://{p}' if p else ''


@login_required
def entrega_pdf(request, pk):
    _gate(request)
    e = get_object_or_404(Entrega.objects.select_related('insumo', 'maquinaria', 'vehiculo', 'registrado_por'), pk=pk)
    km_ant = services.kilometraje_anterior(e.maquinaria, excluir_pk=e.pk) if (e.kilometraje and e.maquinaria_id) else None
    ctx = {'e': e, 'insumo': e.insumo, 'logo': _logo_file(), 'auditora': FIRMA_AUDITORA, 'km_anterior': km_ant,
           'comentarios': e.comentarios.select_related('autor'),
           'recorrido': (e.kilometraje - km_ant) if (km_ant is not None and e.kilometraje is not None) else None}
    return _pdf(request, 'mantenimiento_mecanico/pdf/entrega.html', ctx, f'{e.codigo}.pdf')


def _datos_tabla(insumo, desde, hasta):
    qs = (Entrega.objects.filter(insumo=insumo, fecha__range=(desde, hasta))
          .select_related('maquinaria', 'vehiculo').order_by('fecha', 'hora', 'id'))
    total = qs.aggregate(t=Sum('galones'))['t'] or Decimal('0')
    por_maq = _por_equipo(qs)
    return qs, total, por_maq


def _responsable_firma(qs):
    """<--- hecho por claude code (21-oct-2026): Responsable que FIRMA el reporte = el de las entregas (no quien lo genera).
    Si todas las entregas tienen el mismo responsable, ese; si hay varios o ninguno, línea en blanco (firma a mano)."""
    nombres = {(x or '').strip() for x in qs.values_list('firma_responsable', flat=True) if (x or '').strip()}
    return nombres.pop() if len(nombres) == 1 else ''


@login_required
def reporte_tabla(request, clave, fmt):
    """Reporte TABLA de un insumo en el rango: todas las entregas + totales por maquinaria."""
    _gate(request)
    insumo = _insumo(clave)
    desde, hasta = _rango(request)
    qs, total, por_maq = _datos_tabla(insumo, desde, hasta)
    if fmt == 'pdf':
        ctx = {'insumo': insumo, 'entregas': qs, 'total': total, 'por_maquinaria': por_maq, 'desde': desde, 'hasta': hasta,
               'logo': _logo_file(), 'resumen': services.resumen(insumo), 'auditora': FIRMA_AUDITORA,
               'generado_por': request.user.get_full_name() or request.user.username, 'ahora': datetime.datetime.now(),
               'responsable': _responsable_firma(qs)}
        return _pdf(request, 'mantenimiento_mecanico/pdf/reporte_tabla.html', ctx,
                    f'{insumo.clave}_{desde}_{hasta}.pdf')
    return _excel_tabla(insumo, qs, total, por_maq, desde, hasta)


@login_required
def reporte_general(request, fmt):
    """Reporte GENERAL: todos los insumos con saldo, entradas y entregas del rango, y totales por maquinaria."""
    _gate(request)
    desde, hasta = _rango(request)
    filas = []
    for ins in Insumo.objects.filter(activo=True):
        r = services.resumen(ins)
        ent = Entrada.objects.filter(insumo=ins, fecha__range=(desde, hasta)).aggregate(t=Sum('cantidad'), n=Count('id'))
        sal = Entrega.objects.filter(insumo=ins, fecha__range=(desde, hasta)).aggregate(t=Sum('galones'), n=Count('id'))
        por_maq = _por_equipo(Entrega.objects.filter(insumo=ins, fecha__range=(desde, hasta)))
        filas.append({'insumo': ins, 'resumen': r, 'entradas': ent['t'] or Decimal('0'), 'n_entradas': ent['n'],
                      'entregas': sal['t'] or Decimal('0'), 'n_entregas': sal['n'], 'por_maquinaria': list(por_maq)})
    if fmt == 'pdf':
        ctx = {'filas': filas, 'desde': desde, 'hasta': hasta, 'logo': _logo_file(), 'auditora': FIRMA_AUDITORA,
               'generado_por': request.user.get_full_name() or request.user.username, 'ahora': datetime.datetime.now(),
               'responsable': _responsable_firma(Entrega.objects.filter(fecha__range=(desde, hasta)))}
        return _pdf(request, 'mantenimiento_mecanico/pdf/reporte_general.html', ctx, f'general_{desde}_{hasta}.pdf')
    return _excel_general(filas, desde, hasta)


# ── Excel ────────────────────────────────────────────────────────────────────
def _wb():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    est = {'h': Font(bold=True, color='FFFFFF'), 'hf': PatternFill('solid', fgColor='1E3A8A'),
           'b': Font(bold=True), 'c': Alignment(horizontal='center')}
    return wb, est


def _encabezado(ws, fila, cols, est):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=fila, column=i, value=c)
        cell.font, cell.fill, cell.alignment = est['h'], est['hf'], est['c']


def _xlsx_response(wb, filename):
    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(resp)
    return resp


def _excel_tabla(insumo, qs, total, por_maq, desde, hasta):
    wb, est = _wb()
    ws = wb.active
    ws.title = insumo.nombre[:30]
    ws['A1'] = f'Mantenimiento Mecánico · {insumo.nombre}'
    ws['A1'].font = est['b']
    ws['A2'] = f'Del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}'
    cols = ['Nº', 'Fecha', 'Hora', 'Maquinaria'] + (['Kilometraje'] if insumo.requiere_kilometraje else []) + \
           ['Galones', 'Persona', 'Observaciones', 'Responsable', 'Auditora']
    _encabezado(ws, 4, cols, est)
    f = 5
    for e in qs:
        row = [e.codigo, e.fecha, e.hora.strftime('%H:%M') if e.hora else '', e.equipo_nombre]
        if insumo.requiere_kilometraje:
            row.append(e.kilometraje)
        row += [float(e.galones), e.persona, e.observaciones, e.firma_responsable, e.firma_auditora]
        for i, v in enumerate(row, 1):
            ws.cell(row=f, column=i, value=v)
        ws.cell(row=f, column=2).number_format = 'DD/MM/YYYY'
        f += 1
    ws.cell(row=f, column=1, value='TOTAL').font = est['b']
    ws.cell(row=f, column=cols.index('Galones') + 1, value=float(total)).font = est['b']
    f += 2
    _encabezado(ws, f, ['Maquinaria', 'Entregas', 'Galones'], est)
    for m in por_maq:
        f += 1
        ws.cell(row=f, column=1, value=m['maquinaria__nombre'] or '— sin maquinaria —')
        ws.cell(row=f, column=2, value=m['n'])
        ws.cell(row=f, column=3, value=float(m['gal']))
    for col, w in zip('ABCDEFGHIJ', (12, 12, 8, 26, 12, 10, 24, 36, 22, 16)):
        ws.column_dimensions[col].width = w
    return _xlsx_response(wb, f'{insumo.clave}_{desde}_{hasta}.xlsx')


def _excel_general(filas, desde, hasta):
    wb, est = _wb()
    ws = wb.active
    ws.title = 'Resumen'
    ws['A1'] = 'Mantenimiento Mecánico · Reporte general'
    ws['A1'].font = est['b']
    ws['A2'] = f'Del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}'
    _encabezado(ws, 4, ['Insumo', 'Presentación', 'Entradas (gal)', 'Nº entradas', 'Entregas (gal)', 'Nº entregas',
                        'Saldo actual (gal)', 'Capacidad / barril', 'Mínimo', 'Estado'], est)
    f = 5
    for x in filas:
        ins, r = x['insumo'], x['resumen']
        vals = [ins.nombre, ins.get_presentacion_display(), float(x['entradas']), x['n_entradas'], float(x['entregas']),
                x['n_entregas'], float(r['saldo']), float(ins.capacidad_envase or 0), float(ins.minimo_alerta or 0),
                'SALDO BAJO' if r['bajo'] else 'OK']
        for i, v in enumerate(vals, 1):
            ws.cell(row=f, column=i, value=v)
        f += 1
    for col, w in zip('ABCDEFGHIJ', (22, 22, 14, 11, 14, 11, 16, 16, 10, 12)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet('Por maquinaria')
    _encabezado(ws2, 1, ['Insumo', 'Maquinaria', 'Entregas', 'Galones'], est)
    f = 2
    for x in filas:
        for m in x['por_maquinaria']:
            ws2.cell(row=f, column=1, value=x['insumo'].nombre)
            ws2.cell(row=f, column=2, value=m['maquinaria__nombre'] or '— sin maquinaria —')
            ws2.cell(row=f, column=3, value=m['n'])
            ws2.cell(row=f, column=4, value=float(m['gal']))
            f += 1
    for col, w in zip('ABCD', (22, 30, 10, 12)):
        ws2.column_dimensions[col].width = w
    return _xlsx_response(wb, f'mantenimiento_mecanico_general_{desde}_{hasta}.xlsx')
